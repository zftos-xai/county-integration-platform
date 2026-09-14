from collections import Counter
from copy import copy
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_ROOT = Path(__file__).resolve().parents[2]
INTERFACE_DOCS = PROJECT_ROOT / "docs" / "reference" / "interfaces"
SOURCE = INTERFACE_DOCS / "接口列表.xlsx"
OUTPUT = INTERFACE_DOCS / "县级人民医院_基层和健康V1.0接口矩阵.xlsx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "E7E6E6"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")


def forward_rows(ws):
    system = category = None
    rows = []
    for r in range(2, ws.max_row + 1):
        a, b, name, code = [ws.cell(r, c).value for c in range(1, 5)]
        if a:
            system = str(a).strip()
        if b and str(b).strip() not in {"分类"}:
            category = str(b).strip()
        if not name or str(name).strip() == "接口名称":
            continue
        rows.append({
            "source_row": r,
            "system": system,
            "category": category,
            "name": str(name).strip(),
            "code": str(code).strip() if code else "",
            "method": ws.cell(r, 5).value or "POST（SOAP）",
            "path": ws.cell(r, 6).value or "/WebService.asmx?op=PHIS_Interface",
            "remark": ws.cell(r, 7).value or "",
            "old_priority": ws.cell(r, 8).value,
            "old_implementation": ws.cell(r, 9).value or "",
            "old_change": ws.cell(r, 10).value or "",
        })
    return rows


def proposed_priority(row):
    code, system, name = row["code"], row["system"], row["name"]
    if system == "健康云" and code in {
        "ACB666576D94A00B", "ACB666576D54A002", "ACB666576D94A00D"
    }:
        return "P0-前置"
    if system == "基卫" and code.startswith("100-"):
        return "P0-前置"
    if code == "400-003":
        return "P2-必补"
    if row["remark"] == "不做":
        return "范围外/待决策"
    return {1: "P1", 2: "P2", 3: "P3", 99: "P4-暂缓"}.get(row["old_priority"], "待确认")


def direction(row):
    system, category, code = row["system"], row["category"], row["code"]
    if system == "健康云":
        return "院内集成平台 -> 健康档案云平台"
    if code == "100-009":
        return "HIS 打开第三方页面；第三方回调院内接口验 token"
    if category == "电子病历":
        return "HIS 打开 EMR；EMR -> 院内集成平台 -> HIS"
    if category == "报告跨机构调阅":
        return "HIS -> 互认平台查询/报告页面"
    return "第三方系统/省统建应用 -> 院内集成平台 -> HIS"


def source_and_trigger(row):
    system, category, name, code = row["system"], row["category"], row["name"], row["code"]
    if system == "健康云":
        if name.startswith("查询") or name == "登录验证":
            return "云平台主数据/档案索引", "登录或业务界面按需调用；基础数据定时缓存"
        if name.startswith("上传"):
            return "HIS 业务表/只读视图", "业务完成并审核后写入出站箱异步上传"
        return "HIS 业务表/平台 openId 映射表", "源数据变更后异步修改"
    if code.startswith("100-"):
        return "HIS 主数据与用户权限", "初始化全量同步，后续按日增量；登录按需"
    if code.startswith("200-"):
        return "HIS 挂号、收费和账单", "移动端实时调用；账单按日对账"
    if code.startswith("300-"):
        return "HIS 门诊/住院就诊数据", "转诊或会诊业务按需查询/回写"
    if code.startswith("400-"):
        return "HIS/EMR 病历、诊断、护理数据", "打开病历按需查询；保存/签名后回写"
    if code[:4] in {"500-", "600-", "700-"}:
        return "HIS 申请单与报告表", "医技站扫码拉取；报告审核后回写"
    if code.startswith("800-"):
        return "HIS 挂号、药房、医技队列", "2-5 秒轮询；签到事件实时回写"
    if code.startswith("900-"):
        return "远程会诊系统与 HIS", "会诊报告审核后回写"
    return "互认平台", "诊间按患者身份按需查询"


def dependency(row):
    system, category, code, name = row["system"], row["category"], row["code"], row["name"]
    if system == "健康云":
        if name == "登录验证":
            return "授权审批、c_token、hie_app_key、账号、IP 白名单、NTP"
        if name.startswith("修改"):
            return "登录 token；上传返回 id/sjlyId；区划/用户/字典映射"
        return "登录 token；区划/用户/字典映射；签名参数"
    if code == "100-001":
        return "厂商编号、网络和 WebService 地址"
    if code == "100-002":
        return "厂商编号、HIS 用户、验证码发放规则"
    if code.startswith("100-"):
        return "100-002 验证码；机构/人员权限"
    if code.startswith("200-"):
        return "100-002/003/008；虚拟收费员；收款账户；支付回调"
    if code.startswith("300-"):
        return "100-002/008；患者主索引；FHIR 数据集（300-003）"
    if code.startswith("400-"):
        return "100-009 单点登录；业务ID；医生/科室映射"
    if code[:4] in {"500-", "600-", "700-"}:
        return "100-002/008；申请单唯一号；FHIR 报告规范；项目映射"
    if code.startswith("800-"):
        return "100-002/003；机构、科室和业务ID映射"
    if code.startswith("900-"):
        return "300-001/002；会诊编码唯一约束"
    return "互认平台地址、访问授权、患者身份匹配"


def reliability(row):
    name, code, method = row["name"], row["code"], str(row["method"])
    query = name.startswith(("查询", "获取")) or code.startswith("100-") and code != "100-009"
    if query or method == "GET":
        return "只读；requestId 追踪", "网络/5xx 最多重试2次，指数退避；4xx不重试", "连接3秒，读取15秒；大报文60秒"
    if code.startswith("200-"):
        return "厂商唯一标识+业务单号；先查交易状态再补偿", "禁止盲重试；未知结果先查200-009/012/015", "连接3秒，业务30秒"
    if name.startswith("修改"):
        return "平台 id+sjlyId+源版本号", "出站箱串行重试3次；冲突转人工", "连接3秒，读取20秒"
    if name.startswith("上传"):
        return "机构+源主键+业务类型唯一", "出站箱重试3次；重复先查映射表", "连接3秒，读取20秒"
    if "回写" in name or code.startswith("900-"):
        return "申请单/业务ID+报告类型+版本唯一", "先查本地接收记录；最多重试3次", "连接3秒，业务30秒"
    return "业务主键唯一", "失败入补偿队列，人工可重放", "连接3秒，读取20秒"


def risk(row):
    code, system, name = row["code"], row["system"], row["name"]
    risks = []
    if row["category"] == "报告跨机构调阅":
        return "无独立交易码；接口地址、鉴权、返回字段和页面会话安全定义不足"
    if system == "健康云":
        risks.append("成功码0/200、accessToken大小写及总览/明细路径不一致")
        if name.startswith("修改"):
            risks.append("缺失上传回执将无法修改")
        if name.startswith("上传") or name.startswith("修改"):
            risks.append("必填字段多、位运算字典和日期格式易错")
    else:
        risks.append("HTTP/HTTPS、PHIS_Interface/HIS_Interface、Result/Msg 类型不一致")
        if code == "400-003":
            risks.append("Excel 标记不做，但规范明确要求病历回写")
        if code == "500-003":
            risks.append("详细定义交易编号误写500-004")
        if code == "700-003":
            risks.append("详细定义交易编号误写700-004")
        if code == "300-003":
            risks.append("Result 成功值与公共规范相反")
        if code == "100-009":
            risks.append("指引误写9000，参数表又误写3102")
        if code == "900-001":
            risks.append("参数表误写交易编号1002")
    return "；".join(risks)


def acceptance(row):
    name = row["name"]
    if name.startswith(("查询", "获取")):
        return "必填/空结果/分页/权限/超时通过；关键字段与HIS抽样一致"
    if "登录" in name:
        return "正确/错误/过期凭证、越权、重复登录和审计日志通过"
    return "正常、重复、超时未知、失败重放、回执落库和对账均通过"


def owner(row):
    if row["system"] == "健康云":
        return "集成平台厂商"
    if row["category"] in {"电子病历", "报告跨机构调阅"}:
        return "HIS/EMR厂商"
    return "HIS厂商"


def blocker_ids(row):
    code, system = row["code"], row["system"]
    ids = ["D01", "D02", "D03"] if system == "健康云" else ["D05", "D06"]
    ids.extend({
        "500-003": ["D07"], "700-003": ["D08"], "300-003": ["D09"],
        "100-009": ["D10"], "900-001": ["D11"], "400-003": ["D12"]
    }.get(code, []))
    if row["category"] == "报告跨机构调阅":
        ids.append("D13")
    return "、".join(ids)


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_table(ws, freeze="A2", widths=None):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    style_header(ws)
    if widths:
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width


src = load_workbook(SOURCE, data_only=False)
rows = forward_rows(src["Sheet1"])

# 补充规范中存在但原 Excel 未单列的无交易码能力。
rows.append({
    "source_row": "规范新增", "system": "基卫", "category": "报告跨机构调阅",
    "name": "检查检验报告跨机构调阅", "code": "无独立交易码", "method": "页面/API",
    "path": "由互认平台提供", "remark": "Excel缺项，需纳入范围确认", "old_priority": None,
    "old_implementation": "", "old_change": "HIS工作站增加有/无报告提示与调阅入口"
})

wb = Workbook()
overview = wb.active
overview.title = "总览"

health = [r for r in rows if r["system"] == "健康云"]
phis = [r for r in rows if r["system"] == "基卫" and r["source_row"] != "规范新增"]
unique_phis = {r["code"] for r in phis}

overview_rows = [
    ["县级人民医院对接基层和健康 V1.0 接口矩阵", ""],
    ["统计口径", "结果"],
    ["Excel 健康云接口行数", "=COUNTIF('接口矩阵'!B:B,\"健康云\")"],
    ["Excel 基层 HIS 接口行数", "=COUNTIF('接口矩阵'!B:B,\"基卫\")-1"],
    ["矩阵业务能力总数", "=COUNTA('接口矩阵'!A:A)-1"],
    ["待关闭阻断/高风险项", "=COUNTIF('待确认事项'!B:B,\"阻断\")+COUNTIF('待确认事项'!B:B,\"高\")"],
    ["规范另含无交易码能力", "检查检验报告跨机构调阅"],
    ["通知声明口径", "15类73个：健康云21、基层HIS52；与Excel 74行不一致，需主管部门书面确认"],
    ["推荐架构", "C/S工作站不直连外部平台；统一通过院内集成平台完成协议适配、鉴权、映射、队列、审计和对账"],
    ["首要范围修正", "基础认证/目录接口调整为P0；400-003病历回写调整为P2必补；报告跨机构调阅补入范围"],
    ["生产安全基线", "专网或白名单、TLS、等保要求、最小权限、敏感日志脱敏、凭证集中保管、NTP校时"],
]
for row in overview_rows:
    overview.append(row)
overview.merge_cells("A1:B1")
overview["A1"].fill = PatternFill("solid", fgColor=BLUE)
overview["A1"].font = Font(color=WHITE, bold=True, size=16)
overview["A1"].alignment = Alignment(horizontal="center")
for c in overview[2]:
    c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    c.font = Font(bold=True)
for row in overview.iter_rows(min_row=2):
    for c in row:
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        c.alignment = Alignment(vertical="top", wrap_text=True)
overview.column_dimensions["A"].width = 28
overview.column_dimensions["B"].width = 100
overview.row_dimensions[1].height = 30

matrix = wb.create_sheet("接口矩阵")
headers = [
    "序号", "系统", "业务分类", "接口名称", "事件码/交易码", "协议/方法", "规范地址",
    "原备注", "原优先级", "建议优先级", "调用方向", "数据来源", "触发方式",
    "前置依赖", "幂等键/去重", "重试与补偿", "建议超时", "主要风险", "验收要点",
    "原实现方式", "原功能改造", "来源行", "执行责任方", "当前状态", "阻断编号", "验收证据位置"
]
matrix.append(headers)
for idx, row in enumerate(rows, 1):
    data_source, trigger = source_and_trigger(row)
    idem, retry, timeout = reliability(row)
    matrix.append([
        idx, row["system"], row["category"], row["name"], row["code"], row["method"], row["path"],
        row["remark"], row["old_priority"], proposed_priority(row), direction(row), data_source, trigger,
        dependency(row), idem, retry, timeout, risk(row), acceptance(row), row["old_implementation"],
        row["old_change"], row["source_row"], owner(row), "未开始", blocker_ids(row), "待填写"
    ])
for r in range(2, matrix.max_row + 1):
    priority = matrix.cell(r, 10).value
    color = GREEN if priority in {"P0-前置", "P1"} else YELLOW if priority in {"P2-必补", "P2", "P3"} else RED
    matrix.cell(r, 10).fill = PatternFill("solid", fgColor=color)
style_table(matrix, widths=[6, 10, 18, 28, 22, 16, 30, 20, 10, 14, 38, 25, 32, 36, 30, 38, 26, 48, 42, 38, 38, 12, 18, 16, 18, 34])
matrix.row_dimensions[1].height = 36
status_dv = DataValidation(type="list", formula1='"未开始,合同阻断,环境阻断,开发中,联调中,待验收,已通过,暂缓"', allow_blank=False)
matrix.add_data_validation(status_dv)
status_dv.add(f"X2:X{matrix.max_row}")
for r in range(2, matrix.max_row + 1):
    for c in (23, 24, 26):
        matrix.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)

issues = wb.create_sheet("待确认事项")
issues.append(["编号", "级别", "规范位置", "问题", "建议联调口径/处理", "当前状态", "责任方", "证据/结论"])
issue_rows = [
    (1, "阻断", "健康云5.3、各接口示例", "成功码说明为200，接口示例普遍为0", "联调确认；适配器临时将0/200配置为成功集合，业务回执仍须校验"),
    (2, "阻断", "健康云5.2.1、7.1.1", "请求头写accesstoken，正文又写accessToken", "HTTP头按不区分大小写发送；网关统一规范名并回归测试"),
    (3, "高", "健康云6.3、7章", "总览聚合地址与详细接口实际地址不一致", "以联调环境实际路由为准，禁止硬编码总览短路径"),
    (4, "高", "健康云5.1、7章", "规范要求HTTPS，测试地址为HTTP", "测试网隔离；生产必须专网+TLS并校验证书"),
    (5, "阻断", "基层HIS4.1/4.2/4.5", "HTTPS要求与HTTP地址冲突；PHIS_Interface与HIS_Interface冲突", "由平台主管部门提供WSDL、生产URL、SOAPAction和函数名基线"),
    (6, "高", "基层HIS4.4及各示例", "Result/result、Msg/msg及数字/字符串类型混用，Msg又可能是对象、数组或JSON字符串", "容错解析后归一为统一响应模型；保留原始报文"),
    (7, "阻断", "基层HIS 6.5.4", "接口名500-003，详细交易编号写500-004", "联调书面确认后配置，不在代码中猜测"),
    (8, "阻断", "基层HIS 6.7.4", "接口名700-003，详细交易编号写700-004", "联调书面确认后配置"),
    (9, "阻断", "基层HIS 6.3.3", "300-003参数表写Result 0成功/1失败，与公共规范和示例相反", "以业务回执内容+确认后的成功值判断"),
    (10, "高", "基层HIS 6.1.9", "100-009指引误写9000，参数表又写3102", "固定使用经WSDL联调验证的100-009"),
    (11, "高", "基层HIS 6.9.3", "900-001参数表误写1002", "按总览和标题使用900-001，需书面确认"),
    (12, "高", "基层HIS 6.4.0/6.4.3", "规范提醒病历必须回写，Excel却把400-003标记不做", "纳入P2必做，否则电子病历流程不闭环"),
    (13, "高", "通知与Excel", "通知称73个；Excel实际74行；远程会诊复用两个交易码；跨机构调阅无交易码", "分别按接口行、唯一交易码、业务能力验收并取得范围确认"),
    (14, "高", "两套规范", "未统一规定幂等、超时、重试、熔断、重放和对账", "按本方案在院内集成平台补齐，不让C/S客户端承担可靠性"),
    (15, "高", "基层HIS鉴权", "主要依赖静态验证码，未见签名、时间戳、防重放和限流定义", "专网/IP白名单+TLS+网关限流；验证码加密保存并轮换"),
    (16, "中", "健康云上传/修改", "修改依赖上传返回id、sjlyId，但未规定丢失回执后的恢复查询", "本地回执映射事务落库；未知结果转人工，联调确认补查机制"),
    (17, "中", "基层HIS 6.4.4", "历史病历Content称XML字符串7zip压缩，但未说明传输编码/解压上限", "确认Base64与7z参数；设置解压大小、层级和超时限制"),
    (18, "中", "全规范", "日期格式混用ISO、yyyy-MM-dd HH:mm:ss、/Date(ms)/", "适配层统一内部Instant/LocalDate模型，按接口配置序列化"),
]
issue_rows = [(f"D{number:02d}", *values) for number, *values in issue_rows]
for row in issue_rows:
    issues.append([*row, "待确认", "平台主管部门/医院信息科", "待填写"])
for r in range(2, issues.max_row + 1):
    fill = RED if issues.cell(r, 2).value == "阻断" else YELLOW if issues.cell(r, 2).value == "高" else LIGHT_BLUE
    issues.cell(r, 2).fill = PatternFill("solid", fgColor=fill)
style_table(issues, widths=[8, 10, 24, 70, 70, 14, 28, 42])
issue_status_dv = DataValidation(type="list", formula1='"待确认,处理中,已关闭,接受风险,暂缓"', allow_blank=False)
issues.add_data_validation(issue_status_dv)
issue_status_dv.add(f"F2:F{issues.max_row}")
for r in range(2, issues.max_row + 1):
    for c in (6, 7, 8):
        issues.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)

plan = wb.create_sheet("实施计划")
plan.append(["阶段", "目标", "范围", "关键交付", "准入条件", "退出条件"])
plan_rows = [
    ("P0", "打通底座", "网络、授权、TLS、NTP、日志；云平台登录/区划/用户；基层100-*", "网关、适配器、凭证库、字典/主索引、联调基线", "接入审批完成", "连续运行7天，认证和目录同步成功率>=99.9%"),
    ("P1", "优先业务闭环", "健康档案上传/修改、健康体检上传/修改；PACS/LIS/心电", "视图、出站箱、申请单拉取、FHIR报告回写、对账台账", "P0通过", "重复/超时/重放测试通过，业务抽样一致率100%"),
    ("P2", "电子病历闭环", "400-001~010，必须含400-003", "SSO、病历查询与回写、护理/体温/手术数据", "EMR厂商接口就绪", "病历保存、回写、历史调阅和权限审计通过"),
    ("P3", "协同业务", "档案摘要查询、双向转诊", "工作站入口、患者匹配、FHIR转诊回执", "主索引与机构映射稳定", "转诊全链路及异常补偿通过"),
    ("P4", "按政策启用", "中医保健、高血压、糖尿病、家医签约、便民、叫号、远程、跨机构调阅", "按主管部门确认的范围实施", "业务部门确认责任边界", "分域验收"),
]
for row in plan_rows:
    plan.append(row)
style_table(plan, widths=[10, 22, 55, 60, 35, 55])

mapping = wb.create_sheet("数据治理清单")
mapping.append(["主题", "主键/映射", "治理要求", "责任方", "验收"])
mapping_rows = [
    ("患者主索引", "身份证/其他证件+机构患者ID+健康档案ID", "证件校验、重复合并需人工审批；临时证件按规范生成", "医院信息科+业务科室", "重复率、错配率为0"),
    ("机构", "HIS机构ID、统一社会信用/组织机构代码、平台orgId", "一对一映射并保留生效时间", "信息科", "接口抽样100%一致"),
    ("人员", "HIS职员ID、账号、平台userId/医生ID", "离职停用、科室与执业证书同步", "人事+信息科", "权限与在岗状态一致"),
    ("科室/病区/床位", "本地ID与平台ID", "禁用不删除，维护历史版本", "信息科", "在用目录完整"),
    ("诊断", "ICD-10/中医诊断版本", "明确病案版与医保国临版，禁止只按名称匹配", "医务+病案", "编码映射抽样100%"),
    ("药品/诊疗/耗材", "本地编码、国家医保编码", "全量初始化+时间窗增量；记录版本", "药学/医保/信息科", "数量与行数接口一致"),
    ("公共卫生字典", "性别、民族、婚姻、职业、支付、过敏、疾病等", "位运算字段独立校验；未知值进入隔离队列", "公卫科+信息科", "必填通过率>=99.5%"),
    ("平台回执", "源系统+源主键 -> id/sjlyId/jtId", "与业务状态同事务落库；不可手工覆盖", "集成平台", "所有修改均可追溯上传回执"),
    ("申请单/报告", "申请单ID+业务ID+报告类型+版本", "防重复回写，报告审核后发布，撤销留痕", "医技科室+信息科", "回写与HIS展示一致"),
]
for row in mapping_rows:
    mapping.append(row)
style_table(mapping, widths=[20, 45, 70, 28, 38])

fields = wb.create_sheet("健康档案上传字段")
fields.append(["层级", "序号", "字段名", "字段说明", "必填", "数据类型"])
for row in src["Sheet2"].iter_rows(min_row=2, max_row=src["Sheet2"].max_row, max_col=6, values_only=True):
    fields.append(list(row))
fields.insert_rows(1)
fields["A1"] = "来源：原Excel Sheet2；上传健康档案字段清单。正式开发前应与当前线上Swagger/联调样例再次核对。"
fields.merge_cells("A1:F1")
fields["A1"].fill = PatternFill("solid", fgColor=BLUE)
fields["A1"].font = Font(color=WHITE, bold=True)
fields["A1"].alignment = Alignment(wrap_text=True)
for c in fields[2]:
    c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in fields.iter_rows(min_row=2):
    for c in row:
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        c.alignment = Alignment(vertical="top", wrap_text=True)
fields.freeze_panes = "C2"
for col, width in {"A": 4, "B": 8, "C": 22, "D": 100, "E": 12, "F": 24}.items():
    fields.column_dimensions[col].width = width

readiness = wb.create_sheet("联调准备清单")
readiness.append(["编号", "类别", "前置任务", "所需证据", "执行责任方", "验收责任方", "当前状态", "计划完成日", "证据位置/备注"])
readiness_rows = [
    ("G01", "范围", "冻结机构、业务域、接口代码与使用期限", "批准的接入申请表和范围决议", "医院信息科", "项目负责人", "未开始", "", ""),
    ("G02", "契约", "取得正式WSDL/OpenAPI与版本信息", "原始契约文件、发布日期或哈希", "平台主管部门", "医院信息科", "合同阻断", "", ""),
    ("G03", "契约", "关闭D01-D18中的阻断性冲突", "书面勘误、邮件确认或联调记录", "平台主管部门", "医院信息科", "合同阻断", "", ""),
    ("G04", "网络", "确认测试/生产路由、DNS、端口和出口IP", "网络拓扑、策略单、连通性记录", "医院网络安全", "医院信息科", "未开始", "", ""),
    ("G05", "安全", "验证TLS证书链、白名单和NTP", "证书检查、白名单回执、时钟偏差记录", "医院网络安全", "医院信息科", "未开始", "", ""),
    ("G06", "身份", "发放并登记厂商编号、账号、验证码和应用凭证", "发放记录、保管人与轮换周期，不登记明文密钥", "平台主管部门", "医院信息科", "未开始", "", ""),
    ("G07", "数据", "准备脱敏样例和字段映射", "患者、就诊、申请单、报告、病历样例包", "业务科室/HIS厂商", "医院信息科", "未开始", "", ""),
    ("G08", "数据", "确认机构、人员、科室、诊断和三大目录版本", "映射表、字典版本与差异清单", "HIS厂商", "业务科室", "未开始", "", ""),
    ("G09", "验收", "冻结契约测试与端到端场景", "用例、预期断言、签字人和证据目录", "集成平台厂商", "业务科室/信息科", "未开始", "", ""),
    ("G10", "运维", "准备监控、告警、重放、对账和回退", "演练记录、值班表、告警截图和回退步骤", "集成平台厂商", "医院信息科", "未开始", "", ""),
]
for row in readiness_rows:
    readiness.append(row)
style_table(readiness, widths=[9, 12, 42, 52, 24, 24, 16, 16, 45])
ready_status_dv = DataValidation(type="list", formula1='"未开始,合同阻断,环境阻断,进行中,待验收,已通过,暂缓"', allow_blank=False)
readiness.add_data_validation(ready_status_dv)
ready_status_dv.add(f"G2:G{readiness.max_row}")
for r in range(2, readiness.max_row + 1):
    for c in (5, 6, 7, 8, 9):
        readiness.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)

env = wb.create_sheet("环境参数")
env.append(["环境", "目标系统", "参数项", "要求", "当前值/证据", "责任方", "状态", "备注"])
env_rows = [
    ("测试", "健康云", "基础URL/OpenAPI", "与获批区域和机构一致；记录版本", "待提供", "平台主管部门", "未确认", ""),
    ("测试", "健康云", "出口IP/白名单", "记录申请单与生效时段", "待提供", "医院网络安全", "未确认", ""),
    ("测试", "健康云", "应用凭证", "仅登记凭证编号、保管人和到期日，不填明文", "待提供", "医院信息科", "未确认", ""),
    ("测试", "基层HIS", "WebService URL/WSDL/SOAPAction", "明确HTTPS、端口和PHIS_Interface/HIS_Interface", "待提供", "平台主管部门", "未确认", ""),
    ("测试", "基层HIS", "厂商编号/验证码", "仅登记发放记录，不填明文", "待提供", "平台主管部门", "未确认", ""),
    ("生产", "健康云", "基础URL/证书链/限流", "不得直接沿用测试配置", "待提供", "平台主管部门", "未确认", ""),
    ("生产", "基层HIS", "WebService URL/WSDL/证书链", "与批准的生产基线一致", "待提供", "平台主管部门", "未确认", ""),
    ("全部", "院内集成平台", "NTP/时区", "Asia/Shanghai；记录时钟偏差", "待验证", "医院运维", "未确认", ""),
    ("全部", "院内集成平台", "日志与证据目录", "环境隔离、脱敏、权限和保留期明确", "待提供", "集成平台厂商", "未确认", ""),
]
for row in env_rows:
    env.append(row)
style_table(env, widths=[10, 16, 34, 55, 45, 24, 14, 40])
env_status_dv = DataValidation(type="list", formula1='"未确认,已申请,已提供,已验证,已过期,不适用"', allow_blank=False)
env.add_data_validation(env_status_dv)
env_status_dv.add(f"G2:G{env.max_row}")
for r in range(2, env.max_row + 1):
    for c in (5, 6, 7, 8):
        env.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)

milestones = wb.create_sheet("责任与里程碑")
milestones.append(["里程碑", "任务", "执行责任方", "验收责任方", "前置条件", "当前状态", "计划日期", "完成证据"])
milestone_rows = [
    ("M0", "范围冻结", "医院信息科", "项目负责人", "G01", "未开始", "", "申请表/范围决议"),
    ("M1", "契约基线", "平台主管部门+集成平台厂商", "医院信息科", "G02-G03", "合同阻断", "", "契约、勘误单、版本指纹"),
    ("M2", "环境就绪", "医院网络安全+信息科", "项目负责人", "G04-G06", "未开始", "", "连通性、TLS、白名单、NTP记录"),
    ("M3", "P0单接口通过", "HIS/集成平台厂商", "医院信息科", "M1-M2、G07-G09", "未开始", "", "请求响应、断言、traceId、缺陷记录"),
    ("M4", "P1业务域验收", "HIS/集成平台厂商", "业务科室+信息科", "M3", "未开始", "", "端到端、补偿和对账报告"),
    ("M5", "生产准入", "医院信息科", "变更审批责任人", "M4、G10", "未开始", "", "上线检查、回退演练、签字记录"),
]
for row in milestone_rows:
    milestones.append(row)
style_table(milestones, widths=[12, 28, 30, 26, 24, 16, 16, 52])
milestone_status_dv = DataValidation(type="list", formula1='"未开始,合同阻断,环境阻断,进行中,待验收,已通过,暂缓"', allow_blank=False)
milestones.add_data_validation(milestone_status_dv)
milestone_status_dv.add(f"F2:F{milestones.max_row}")
for r in range(2, milestones.max_row + 1):
    for c in (3, 4, 6, 7, 8):
        milestones.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)

cutover = wb.create_sheet("上线检查表")
cutover.append(["编号", "类别", "检查项", "通过标准", "责任方", "状态", "证据位置", "回退/停写触发"])
cutover_rows = [
    ("C01", "范围", "生产机构与接口开关", "与批准范围一致，未批准域默认关闭", "医院信息科", "未检查", "", "发现越范围调用"),
    ("C02", "数据", "患者、机构、人员和目录抽样", "关键标识100%匹配", "业务科室/HIS厂商", "未检查", "", "患者错配或跨机构数据"),
    ("C03", "可靠性", "重复、超时未知、重放和对账", "测试用例全部通过且状态可追踪", "集成平台厂商", "未检查", "", "无法确认写入结果或无法补偿"),
    ("C04", "性能", "峰值容量与队列清空", "P95和容量达到评审基线", "集成平台厂商", "未检查", "", "P95连续15分钟超过基线2倍或队列预计2小时无法清空"),
    ("C05", "安全", "TLS、白名单、最小权限、脱敏、凭证轮换", "安全检查全部通过", "医院网络安全", "未检查", "", "凭证泄露、未授权访问或敏感日志外泄"),
    ("C06", "监控", "接口、业务、队列、资源和证书告警", "告警接收与升级链验证通过", "医院运维", "未检查", "", "关键状态不可观测或告警失效"),
    ("C07", "回退", "停写、冻结队列、恢复配置和数据核对演练", "演练完成并有签字记录", "医院信息科", "未检查", "", "无法执行批准的回退步骤"),
    ("C08", "值守", "上线窗口、联系人、升级路径和日报", "人员到位且联系方式验证", "项目负责人", "未检查", "", "关键责任人不可用"),
]
for row in cutover_rows:
    cutover.append(row)
style_table(cutover, widths=[9, 12, 40, 52, 24, 14, 42, 55])
cutover_status_dv = DataValidation(type="list", formula1='"未检查,不通过,待复核,已通过,不适用"', allow_blank=False)
cutover.add_data_validation(cutover_status_dv)
cutover_status_dv.add(f"F2:F{cutover.max_row}")
for r in range(2, cutover.max_row + 1):
    for c in (5, 6, 7):
        cutover.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)

changes = wb.create_sheet("变更记录")
changes.append(["版本", "日期", "变更内容", "编制", "审核", "状态"])
changes.append(["V1.1", "2026-09-10", "新增逐接口责任/状态/阻断/证据字段，以及联调准备、环境参数、责任里程碑和上线检查表", "待填写", "待填写", "评审稿"])
style_table(changes, widths=[12, 16, 85, 20, 20, 16])
for c in (4, 5, 6):
    changes.cell(2, c).fill = PatternFill("solid", fgColor=YELLOW)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                font = copy(cell.font)
                font.name = "Arial"
                cell.font = font

wb.save(OUTPUT)
print(f"已生成 {OUTPUT}，接口矩阵 {len(rows)} 行（含规范新增1项）")
